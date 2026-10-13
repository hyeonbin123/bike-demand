"""따릉이 대여소 정보(반기별 xlsx)를 Parquet 하나로 모은다.

파일마다 머리글이 5줄이고 앞 10개 컬럼의 위치가 같다. 값은 문자열로 두고
(해석은 dbt에서), 어느 시점 파일인지를 ``snapshot``(YYYY-MM)으로 붙인다.
대여소번호는 xlsx에서 숫자(301)지만 대여이력에서는 앞자리 0이 붙은 문자열("00301")이다.
"""

# [파일 역할] 서울 열린데이터광장 '공공자전거 대여소 정보'(OA-13252) 반기 xlsx 8개(22.12월 ~ 26.6월
# 기준)를 bronze Parquet 하나로 모은다. 값 해석(숫자·좌표 변환, 거치대 합산)은 dbt stg_stations가
# 한다.
# [왜 시점(snapshot)을 붙이나] 대여소의 이름·위치·거치대 수는 반기마다 바뀐다. dim_stations는 최신
# 시점을 쓰고, 학습 행은 그 달보다 앞선 시점의 값을 써야 미래 정보가 섞이지
# 않는다(docs/experiments.md v3). 그래서 파일 이름의 기준 시점을 행마다 남긴다.
# [연결] refresh_history DAG의 convert_stations 작업이 실행부를 부른다. 결과는 dbt source
# bronze.stations로 읽힌다.

from __future__ import annotations

import datetime as dt
import os
import re
from pathlib import Path

# openpyxl: xlsx 읽기(read_only 모드로 메모리를 아낌). pyarrow: 문자열 컬럼 Parquet 쓰기.
import openpyxl
import pyarrow as pa
import pyarrow.parquet as pq

# HEADER_ROWS: 파일마다 표 위에 머리글이 5줄 있다. COLUMNS: 앞 10개 칸의 프로젝트 컬럼 이름(파일마다
# 위치가 같음). docks_lcd·docks_qr은 거치 방식별 거치대 수로, dbt에서 더해 docks가 된다.
HEADER_ROWS = 5
COLUMNS = [
    "station_no",
    "station_name",
    "district",
    "address",
    "lat",
    "lon",
    "installed_at",
    "docks_lcd",
    "docks_qr",
    "operation_type",
]
# 모든 값을 문자열로 저장하고, 기준 시점(snapshot, YYYY-MM)과 원본 파일 이름(source_file)을
# 덧붙인다.
SCHEMA = pa.schema(
    [(name, pa.string()) for name in COLUMNS]
    + [("snapshot", pa.string()), ("source_file", pa.string())]
)
# 파일 이름의 '(26.6월 기준)' 같은 부분에서 연도 두 자리와 월을 뽑는 정규식.
_SNAPSHOT_RE = re.compile(r"\((\d{2})\.(\d{1,2})월 기준\)")
# 머리글 첫 줄 앞 세 칸. 양식이 바뀌면 조용히 잘못 읽지 않도록 확인한다.
EXPECTED_HEADER = ("대여소\n번호", "보관소(대여소)명", "소재지(위치)")


# 파일 이름에서 기준 시점을 YYYY-MM으로 만든다. 예: '(22.12월 기준)' → 2022-12. 기준 시점을 못
# 찾으면 어느 반기 정보인지 모르므로 멈춘다.
def snapshot_of(filename: str) -> str:
    match = _SNAPSHOT_RE.search(filename)
    if not match:
        raise ValueError(f"파일 이름에서 기준 시점(YY.M월 기준)을 찾지 못함: {filename}")
    yy, mm = match.groups()
    return f"20{yy}-{int(mm):02d}"


# 셀 값을 문자열로 바꾼다. 날짜 셀(설치일)은 'YYYY-MM-DD HH:MM:SS' 형식으로, 빈 칸·공백뿐인 칸은
# None으로 바꾼다.
def _text(value: object) -> str | None:
    if value is None:
        return None
    if isinstance(value, dt.datetime):
        return value.isoformat(sep=" ")
    text = str(value).strip()
    return text or None


# xlsx 한 파일을 행 사전 목록으로 읽는다. data_only=True는 수식 대신 계산된 값을 읽는다. 호출 위치:
# to_parquet.
def read_workbook(path: Path) -> list[dict]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        rows = sheet.iter_rows(values_only=True)
        header = next(rows)
        # 첫 줄 앞 세 칸이 예상한 머리글과 다르면(양식 변경) 칸 위치가 어긋났을 수 있으므로 조용히
        # 잘못 읽지 않고 멈춘다. 확인 뒤 남은 머리글 4줄을 건너뛴다.
        if tuple(header[:3]) != EXPECTED_HEADER:
            raise ValueError(f"{path.name}: 예상과 다른 머리글 {header[:3]}")
        for _ in range(HEADER_ROWS - 1):
            next(rows)
        snapshot = snapshot_of(path.name)
        records = []
        for row in rows:
            values = [_text(v) for v in row[: len(COLUMNS)]]
            # 대여소번호가 빈 줄(표 끝의 빈 행 등)은 건너뛴다.
            if values[0] is None:  # 빈 줄
                continue
            record = dict(zip(COLUMNS, values, strict=True))
            record.update(snapshot=snapshot, source_file=path.name)
            records.append(record)
        return records
    # read_only 모드는 파일을 열어 둔 채 읽으므로 끝나면 반드시 닫는다.
    finally:
        workbook.close()


# raw_dir의 모든 xlsx를 이름순으로 읽어 Parquet 하나를 통째로 다시 만든다. 같은 기준 시점 파일이
# 둘이면 어느 것이 맞는지 모르므로 멈춘다. 임시 파일에 쓰고 os.replace로 바꿔 dbt가 반쯤 쓴 파일을
# 읽지 않게 한다.
# 호출 위치: 아래 실행부, tests/test_ingest_stations.py.
def to_parquet(raw_dir: Path, out_path: Path) -> dict[str, int]:
    """모든 xlsx를 읽어 Parquet 하나로 다시 만들고 시점별 행 수를 돌려준다."""
    records: list[dict] = []
    counts: dict[str, int] = {}
    for path in sorted(raw_dir.glob("*.xlsx")):
        rows = read_workbook(path)
        snapshot = snapshot_of(path.name)
        if snapshot in counts:
            raise ValueError(f"{snapshot} 시점 파일이 두 개")
        counts[snapshot] = len(rows)
        records.extend(rows)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    tmp = out_path.with_suffix(".parquet.tmp")
    pq.write_table(pa.Table.from_pylist(records, schema=SCHEMA), tmp, compression="zstd")
    os.replace(tmp, out_path)
    return dict(sorted(counts.items()))


if __name__ == "__main__":
    import argparse

    # 실행: python -m bike_demand.ingest.stations. 기준 시점별 행 수를 출력해 새 반기 파일이
    # 들어갔는지 눈으로 확인할 수 있게 한다.
    parser = argparse.ArgumentParser(description="대여소 정보 xlsx를 Parquet로 모은다")
    parser.add_argument("--raw", type=Path, default=Path("data/raw/stations"))
    parser.add_argument("--out", type=Path, default=Path("data/bronze/stations/stations.parquet"))
    args = parser.parse_args()
    for snapshot, rows in to_parquet(args.raw, args.out).items():
        print(snapshot, rows)
