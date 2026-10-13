# [무엇을 검증하나] 반기별 대여소 정보 xlsx를 하나의 문자열 Parquet로 모으는 ingest/stations.py를
# 확인한다: 파일 이름의 기준 시점(YY.M월 기준 → YYYY-MM) 읽기, 5줄 머리글 건너뛰기, 값을 해석하지
# 않고 문자열로 두기, 빈 줄 건너뛰기, 머리글 양식이 바뀌면 실패하기.
# [왜 중요한가] 이 스냅샷들은 v3에서 '행보다 앞선 시점의 대여소 정보(거치대 수·좌표·자치구)'를
# 붙이는 근거다(docs/experiments.md v3). 시점이 틀리거나 양식 변경으로 열이 밀리면 과거 행에 엉뚱한
# 대여소 정보가 붙는다.
# [방식] openpyxl로 실제 원본과 같은 모양(머리글 5줄, 첫 시트)의 작은 xlsx를 임시 폴더에 만든다.

import datetime as dt

# openpyxl: 테스트용 xlsx를 만든다. pyarrow.parquet: 결과 Parquet을 읽어 확인한다.
import openpyxl
import pyarrow.parquet as pq
import pytest

from bike_demand.ingest import stations

# 실제 원본의 5줄 머리글을 흉내 낸다. 첫 줄 앞 세 칸은 모듈이 양식 확인에 쓰는
# EXPECTED_HEADER(줄바꿈이 든 셀 이름 포함)이고, 나머지 줄은 병합 셀처럼 군데군데 비어 있다.
HEADER = [
    [*stations.EXPECTED_HEADER, None, None, None, "설치\n시기", "설치형태"],
    [None, None, None, None, None, None, None, "LCD", "QR"],
    [None, None, "자치구", "상세주소", "위도", "경도"],
    [None, None, None, None, None, None, None, "거치\n대수", "거치\n대수"],
    [],
]


# 머리글 + 행으로 xlsx를 만든다. 대여소 시트를 맨 앞에 두고 기본 시트를 지워, 모듈이 읽는 '첫
# 시트'가 자료 시트가 되게 한다.
def make_xlsx(path, rows, header=HEADER):
    workbook = openpyxl.Workbook()
    sheet = workbook.create_sheet("대여소현황", 0)
    del workbook[workbook.sheetnames[1]]
    for row in header + rows:
        sheet.append(row)
    workbook.save(path)


# '(24.6월 기준)' → 2024-06, '(22.12월 기준)' → 2022-12. 기준 시점이 없는 파일 이름은 추측하지 않고
# 오류.
def test_snapshot_of():
    assert stations.snapshot_of("공공자전거 대여소 정보(24.6월 기준).xlsx") == "2024-06"
    assert stations.snapshot_of("공공자전거 대여소 정보(22.12월 기준).xlsx") == "2022-12"
    with pytest.raises(ValueError):
        stations.snapshot_of("대여소.xlsx")


# 두 시점 파일을 모은다. 26.6월 파일의 빈 줄은 건너뛰어 1행, 25.12월 파일은 같은 행 두 줄을 걸러
# 내지 않고 그대로 2행으로 센다.
# 값은 모두 문자열: 숫자 대여소번호 301 → '301', 앞 공백은 제거, 설치 일시 datetime → 'YYYY-MM-DD
# HH:MM:SS', 빈 칸은 None. LCD·QR 거치대 수가 열 위치대로 들어가는지도 본다.
def test_to_parquet_reads_rows_as_text(tmp_path):
    raw = tmp_path / "raw"
    raw.mkdir()
    row = [301, " 경복궁역 7번출구 앞", "종로구", "주소", 37.5757, 126.9714,
           dt.datetime(2015, 10, 7, 12, 3, 46), None, 20, "QR", None]  # fmt: skip
    make_xlsx(raw / "공공자전거 대여소 정보(26.6월 기준).xlsx", [row, [None] * 10])
    make_xlsx(raw / "공공자전거 대여소 정보(25.12월 기준).xlsx", [row, row])

    out = tmp_path / "stations.parquet"
    assert stations.to_parquet(raw, out) == {"2025-12": 2, "2026-06": 1}
    first = [r for r in pq.read_table(out).to_pylist() if r["snapshot"] == "2026-06"][0]
    assert first["station_no"] == "301"
    assert first["station_name"] == "경복궁역 7번출구 앞"  # 앞 공백 제거
    assert first["installed_at"] == "2015-10-07 12:03:46"
    assert first["docks_lcd"] is None and first["docks_qr"] == "20"


# 첫 줄 머리글이 기대와 다르면 '머리글' 오류로 멈춘다. 양식이 바뀐 파일을 열 위치만 믿고 읽으면
# 조용히 다른 열이 들어가기 때문이다(모듈 주석).
def test_unexpected_header_fails(tmp_path):
    bad = [["번호", "이름", "위치"]] + HEADER[1:]
    path = tmp_path / "공공자전거 대여소 정보(26.6월 기준).xlsx"
    make_xlsx(path, [], header=bad)
    with pytest.raises(ValueError, match="머리글"):
        stations.read_workbook(path)
